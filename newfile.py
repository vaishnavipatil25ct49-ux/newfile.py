	from openpyxl import Workbook, load_workbook
import os

FILE_NAME = "student_results.xlsx"


# Calculate total, percentage, grade and status
def calculate_result(marks):

    total = sum(marks)
    percentage = total / 5

    # Grade
    if percentage >= 90:
        grade = "A+"
    elif percentage >= 80:
        grade = "A"
    elif percentage >= 70:
        grade = "B"
    elif percentage >= 60:
        grade = "C"
    elif percentage >= 50:
        grade = "D"
    else:
        grade = "F"

    # Pass / Fail
    if all(mark >= 35 for mark in marks):
        status = "PASS"
    else:
        status = "FAIL"

    return total, percentage, grade, status


# Add student
def add_student():

    roll_no = input("Enter Roll No: ")
    name = input("Enter Student Name: ")
    course = input("Enter Course/Class: ")

    marks = []

    for i in range(1, 6):
        mark = float(input("Enter marks of Subject " + str(i) + ": "))
        marks.append(mark)

    total, percentage, grade, status = calculate_result(marks)

    # Create Excel file if it does not exist
    if not os.path.exists(FILE_NAME):

        workbook = Workbook()
        sheet = workbook.active

        sheet.append([
            "Roll No",
            "Name",
            "Class",
            "Subject 1",
            "Subject 2",
            "Subject 3",
            "Subject 4",
            "Subject 5",
            "Total",
            "Percentage",
            "Grade",
            "Status"
        ])

        workbook.save(FILE_NAME)

    # Open Excel file
    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    # Add student record
    sheet.append([
        roll_no,
        name,
        course,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        grade,
        status
    ])

    workbook.save(FILE_NAME)

    print("\nStudent result saved successfully!")
    print("Total      :", total)
    print("Percentage :", format(percentage, ".2f") + "%")
    print("Grade      :", grade)
    print("Status     :", status)


# Get one student's result
def get_result():

    roll_no = input("Enter Roll No to search: ")

    if not os.path.exists(FILE_NAME):
        print("No student data found.")
        return

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == roll_no:

            print("\n--------------------------------")
            print("        Student Result")
            print("--------------------------------")
            print("Roll No     :", row[0])
            print("Name        :", row[1])
            print("Class       :", row[2])
            print("Total       :", row[8])
            print("Percentage  :", format(row[9], ".2f") + "%")
            print("Grade       :", row[10])
            print("Status      :", row[11])
            print("--------------------------------")

            found = True
            break

    if not found:
        print("Student not found.")


# Show all student data
def show_all_data():

    if not os.path.exists(FILE_NAME):
        print("No student data found.")
        return

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    print("\n---------------------------------------------------------------")
    print("                     ALL STUDENT DATA")
    print("---------------------------------------------------------------")

    print(
        "Roll No\tName\tClass\tTotal\tPercentage\tGrade\tStatus"
    )

    for row in sheet.iter_rows(min_row=2, values_only=True):

        print(
            row[0], "\t",
            row[1], "\t",
            row[2], "\t",
            row[8], "\t",
            format(row[9], ".2f"), "\t\t",
            row[10], "\t",
            row[11]
        )


# Main menu
def menu():

    while True:

        print("\n========================================")
        print("       STUDENT RESULT MANAGEMENT")
        print("========================================")
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")
        print("========================================")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("Thank you!")
            break

        else:
            print("Invalid choice. Please try again.")


# Start program
menu()
        

    
        
			
		